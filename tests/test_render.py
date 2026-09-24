import random
from decimal import ROUND_HALF_UP, Decimal

import pytest
from faker import Faker
from openpyxl import load_workbook

from invoiceops.contract import Invoice, build_contract
from invoiceops.synth.carriers import make_carriers
from invoiceops.synth.render import make_invoice, write_xlsx


def setup(seed):
    rng = random.Random(seed)
    fake = Faker("en_US")
    fake.seed_instance(seed)
    return make_carriers(1, rng, fake)[0], rng, fake


@pytest.mark.parametrize("seed", range(20))
def test_invoice_totals_are_consistent(seed):
    spec, rng, fake = setup(seed)
    rec = make_invoice(spec, "fr", rng, fake, 0)
    assert rec["net"] == sum(line["amount"] for line in rec["lines"])
    assert rec["gross"] == rec["net"] + rec["tax"]
    for line in rec["lines"]:
        assert line["amount"] == (line["quantity"] * line["unit_price"]).quantize(Decimal("0.01"), ROUND_HALF_UP)
    assert isinstance(build_contract(rec), Invoice)


@pytest.mark.parametrize("seed", range(10))
def test_xlsx_header_row_matches_meta(seed, tmp_path):
    spec, rng, fake = setup(seed)
    rec = make_invoice(spec, "es", rng, fake, 1)
    meta = write_xlsx(tmp_path / "x.xlsx", spec, "es", rec, rng, fake)
    rows = [list(r) for r in load_workbook(tmp_path / "x.xlsx").active.iter_rows(values_only=True)]
    assert meta["header_row"] == spec.n_banner_rows
    assert rows[meta["header_row"]] == meta["headers"]
    assert len(rows) == meta["header_row"] + 1 + len(rec["lines"])
    assert len(meta["mapping"]) == len(meta["headers"])
    data_rows = rows[meta["header_row"] + 1 :]
    for col, field in enumerate(meta["mapping"]):
        if field == "amount":
            for row, line in zip(data_rows, rec["lines"]):
                assert Decimal(str(row[col])) == line["amount"]
        elif field == "currency":
            for row in data_rows:
                assert row[col] == spec.currency


@pytest.mark.parametrize("seed", range(200))
def test_invoice_has_at_least_two_charge_types(seed):
    spec, rng, fake = setup(seed)
    rec = make_invoice(spec, "fr", rng, fake, 0)
    assert len({line["charge_type"] for line in rec["lines"]}) >= 2


def test_content_sha256_reproducible_and_sensitive_to_inputs(tmp_path):
    seed = 0
    spec, rng, fake = setup(seed)
    rec = make_invoice(spec, "es", rng, fake, 1)

    rng_a = random.Random(seed)
    fake_a = Faker("en_US")
    fake_a.seed_instance(seed)
    meta_a = write_xlsx(tmp_path / "a.xlsx", spec, "es", rec, rng_a, fake_a)

    rng_b = random.Random(seed)
    fake_b = Faker("en_US")
    fake_b.seed_instance(seed)
    meta_b = write_xlsx(tmp_path / "b.xlsx", spec, "es", rec, rng_b, fake_b)

    assert meta_a["content_sha256"] == meta_b["content_sha256"]

    rng_c = random.Random(seed + 1)
    fake_c = Faker("en_US")
    fake_c.seed_instance(seed + 1)
    meta_c = write_xlsx(tmp_path / "c.xlsx", spec, "es", rec, rng_c, fake_c)

    assert meta_c["content_sha256"] != meta_a["content_sha256"]


def test_xlsx_content_is_reproducible(tmp_path):
    seed = 0
    spec, rng, fake = setup(seed)
    rec = make_invoice(spec, "es", rng, fake, 1)

    rng2 = random.Random(seed)
    fake2 = Faker("en_US")
    fake2.seed_instance(seed)
    meta_a = write_xlsx(tmp_path / "a.xlsx", spec, "es", rec, rng2, fake2)

    rng3 = random.Random(seed)
    fake3 = Faker("en_US")
    fake3.seed_instance(seed)
    meta_b = write_xlsx(tmp_path / "b.xlsx", spec, "es", rec, rng3, fake3)

    assert meta_a == meta_b
    rows_a = [list(r) for r in load_workbook(tmp_path / "a.xlsx").active.iter_rows(values_only=True)]
    rows_b = [list(r) for r in load_workbook(tmp_path / "b.xlsx").active.iter_rows(values_only=True)]
    assert rows_a == rows_b
