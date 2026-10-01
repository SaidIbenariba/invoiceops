"""Read XLSX/CSV into rows and recover the header row below any banner rows."""

import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class Table:
    header_row: int
    headers: list[str]
    rows: list[list[object]]

    def column(self, i: int) -> list[object]:
        return [r[i] if i < len(r) else None for r in self.rows]


def _blank(cell: object) -> bool:
    return cell is None or (isinstance(cell, str) and not cell.strip())


def _is_number(cell: object) -> bool:
    return isinstance(cell, (int, float)) and not isinstance(cell, bool)


def _coerce(cell: str) -> object:
    if not cell.strip():
        return None
    try:
        return float(cell)
    except ValueError:
        return cell


def read_rows(path: Path) -> list[list[object]]:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8") as f:
            return [[_coerce(c) for c in row] for row in csv.reader(f)]
    ws = load_workbook(path, data_only=True).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def find_header_row(rows: list[list[object]], max_scan: int = 15) -> int:
    """First row with >= 3 non-blank cells, all text, followed by a row containing a number."""
    for i, row in enumerate(rows[:max_scan]):
        cells = [c for c in row if not _blank(c)]
        if len(cells) < 3 or not all(isinstance(c, str) for c in cells):
            continue
        following = rows[i + 1] if i + 1 < len(rows) else []
        if any(_is_number(c) for c in following):
            return i
    raise ValueError(f"no header row found in the first {max_scan} rows")


def table_from_rows(rows: list[list[object]]) -> Table:
    h = find_header_row(rows)
    width = len(rows[h])
    headers = ["" if _blank(c) else str(c).strip() for c in rows[h]]
    body = []
    for r in rows[h + 1 :]:
        if all(_blank(c) for c in r):
            break
        body.append(list(r) + [None] * (width - len(r)))
    return Table(header_row=h, headers=headers, rows=body)


def load_table(path: Path) -> Table:
    return table_from_rows(read_rows(path))
